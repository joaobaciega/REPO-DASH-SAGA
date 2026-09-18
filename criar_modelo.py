"""Cria uma cópia da planilha atual com abas vazias para verbas e pagamentos."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from configuracao import BASE_DIR
from planilha import VERBAS_COLS


def main():
    saida = BASE_DIR / "MODELO SAGA COMPLETO.xlsx"
    if saida.exists():
        raise SystemExit("MODELO SAGA COMPLETO.xlsx já existe; ele foi preservado.")
    wb = load_workbook(BASE_DIR / "BASE GRUPO SAGA.xlsx")
    for nome, colunas in (("Verbas", VERBAS_COLS),
                           ("Pagamentos", ["Mês", "Consultor Pago", "Gerente Pago"]),
                           ("Marketing", ["Mês", "Valor"])):
        if nome not in wb.sheetnames:
            wb.create_sheet(nome).append(colunas)
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="16376A")
            cell.font = Font(color="FFFFFF", bold=True)
            ws.column_dimensions[cell.column_letter].width = 23
    wb.save(saida)
    wb.close()
    print(f"Modelo criado: {saida}")


if __name__ == "__main__":
    main()
