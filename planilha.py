"""Leitura e validação da planilha Saga, sem depender de MySQL ou do Streamlit."""
import datetime as dt
import hashlib
import json
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

TABELAS = ("unidades", "consultores", "consultor_unidade", "lancamentos",
           "vendas_verbas", "verbas_pagamentos", "verbas_marketing_pagos")


def identificador(*partes):
    return hashlib.sha256(json.dumps(partes, ensure_ascii=False).encode()).hexdigest()


def texto(valor, contexto, limite=180):
    if valor is None or not str(valor).strip():
        raise ValueError(f"{contexto}: texto obrigatório vazio.")
    resultado = " ".join(str(valor).split())
    if len(resultado) > limite:
        raise ValueError(f"{contexto}: máximo de {limite} caracteres.")
    return resultado


def numero(valor, contexto, *, inteiro=False, opcional=False):
    if valor is None or valor == "":
        if opcional:
            return None
        raise ValueError(f"{contexto}: número obrigatório vazio (ou fórmula sem resultado salvo).")
    if isinstance(valor, (dt.date, dt.datetime, bool)):
        raise ValueError(f"{contexto}: esperado número, recebido {type(valor).__name__}.")
    try:
        s = str(valor).replace("R$", "").replace(" ", "").strip()
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        n = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"{contexto}: número inválido.") from None
    if not n.is_finite() or n < 0 or n > Decimal("9999999999.99"):
        raise ValueError(f"{contexto}: valor deve ser finito, não negativo e menor que 10 bilhões.")
    if inteiro:
        if n != n.to_integral_value() or n > 2147483647:
            raise ValueError(f"{contexto}: quantidade deve ser inteira e menor que 2.147.483.648.")
        return int(n)
    if n != n.quantize(Decimal("0.01")):
        raise ValueError(f"{contexto}: utilize no máximo duas casas decimais.")
    return n.quantize(Decimal("0.01"))


def data(valor, contexto, *, mes=False):
    if isinstance(valor, dt.datetime):
        resultado = valor.date()
    elif isinstance(valor, dt.date):
        resultado = valor
    else:
        resultado = None
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%m/%Y", "%Y-%m"):
            try:
                resultado = dt.datetime.strptime(str(valor).strip(), formato).date()
                break
            except ValueError:
                pass
        if resultado is None:
            raise ValueError(f"{contexto}: data inválida; use uma data do Excel ou DD/MM/AAAA.")
    return resultado.replace(day=1) if mes else resultado


def linhas(wb, aba, obrigatorias):
    if aba not in wb.sheetnames:
        raise ValueError(f"Aba obrigatória ausente: {aba}.")
    valores = iter(wb[aba].values)
    cab = [str(x).strip() if x is not None else "" for x in next(valores)]
    if len([x for x in cab if x]) != len(set(x for x in cab if x)):
        raise ValueError(f"Aba {aba}: cabeçalhos repetidos.")
    faltam = set(obrigatorias) - set(cab)
    if faltam:
        raise ValueError(f"Aba {aba}: faltam as colunas {', '.join(sorted(faltam))}.")
    for linha, vals in enumerate(valores, 2):
        if all(v is None or v == "" for v in vals):
            continue
        yield dict(zip(cab, vals)), f"{aba}, linha {linha}"


@dataclass
class Base:
    tabelas: dict
    avisos: list
    sha256: str
    arquivo: str

    @property
    def faturamento(self):
        return sum((r["refil_diant"] * r["preco_diant"] + r["refil_tras"] * r["preco_tras"]
                    for r in self.tabelas["lancamentos"]), Decimal("0.00"))


def ler_planilha(caminho):
    caminho = Path(caminho)
    # Lê uma única versão: hash e conteúdo não divergem se o Excel for salvo durante a carga.
    import io
    conteudo = caminho.read_bytes()
    wb = load_workbook(io.BytesIO(conteudo), data_only=True)
    tabelas = {t: [] for t in TABELAS}
    avisos = []
    precos, gerentes, resumos = {}, {}, {}
    try:
        for r, ctx in linhas(wb, "Preço por Unidade", ["Loja", "Marca", "Preço Diant. (R$)", "Preço Tras. (R$)"]):
            chave = (texto(r["Marca"], ctx, 140), texto(r["Loja"], ctx, 140))
            if chave in precos:
                raise ValueError(f"{ctx}: preço duplicado para {chave}.")
            precos[chave] = (numero(r["Preço Diant. (R$)"], ctx), numero(r["Preço Tras. (R$)"], ctx))
        for r, ctx in linhas(wb, "Gerentes", ["Gerente", "Loja", "Marca", "Mês (data)"]):
            chave = (texto(r["Marca"], ctx, 140), texto(r["Loja"], ctx, 140), data(r["Mês (data)"], ctx, mes=True))
            if chave in gerentes:
                raise ValueError(f"{ctx}: dois gerentes para a mesma loja/marca/mês.")
            gerentes[chave] = texto(r["Gerente"], ctx)
            resumos[chave] = (r, ctx)
        unis, cons, fatos, vinculos = {}, {}, {}, {}
        cols = ["Consultor", "Loja", "Marca", "Mês (data)", "Passagens", "Refil Diant. (qtd)", "Refil Tras. (qtd)"]
        for r, ctx in linhas(wb, "Consultores", cols):
            nome = texto(r["Consultor"], ctx)
            marca, loja = texto(r["Marca"], ctx, 140), texto(r["Loja"], ctx, 140)
            mes = data(r["Mês (data)"], ctx, mes=True)
            cid, uid = identificador(nome), identificador(marca, loja)
            chave = (cid, uid, mes)
            if chave in fatos:
                raise ValueError(f"{ctx}: consultor/loja/marca/mês duplicado. Informe o acumulado em uma única linha.")
            if (marca, loja) not in precos:
                raise ValueError(f"{ctx}: loja sem cadastro em Preço por Unidade.")
            if (marca, loja, mes) not in gerentes:
                raise ValueError(f"{ctx}: gerente não cadastrado para esta loja/marca/mês.")
            gerente = gerentes[marca, loja, mes]
            pd, pt = precos[marca, loja]
            # Preços explícitos na linha congelam meses antigos após reajuste da tabela atual.
            preco_d = numero(r.get("Preço Diant. (R$)"), ctx, opcional=True)
            preco_t = numero(r.get("Preço Tras. (R$)"), ctx, opcional=True)
            fato = dict(consultor_id=cid, unidade_id=uid, mes=mes, gerente=gerente,
                        passagens=numero(r["Passagens"], ctx, inteiro=True, opcional=True),
                        refil_diant=numero(r["Refil Diant. (qtd)"], ctx, inteiro=True),
                        refil_tras=numero(r["Refil Tras. (qtd)"], ctx, inteiro=True),
                        preco_diant=pd if preco_d is None else preco_d,
                        preco_tras=pt if preco_t is None else preco_t)
            fatos[chave] = fato
            cons[cid] = dict(id=cid, nome=nome)
            unis[uid] = dict(id=uid, nome_exibicao=f"{marca} · {loja}", marca=marca, loja=loja,
                             gerente=gerente, preco_diant=pd, preco_tras=pt)
            vinculos.setdefault(cid, []).append((mes, uid))
            if fato["passagens"] is None:
                avisos.append(f"{ctx}: passagens não informadas; conversão ficará indisponível nesta linha.")
            elif fato["refil_diant"] > fato["passagens"]:
                avisos.append(f"{ctx}: refis dianteiros acima das passagens; confira os valores.")
        if not fatos:
            raise ValueError("Aba Consultores vazia. Substituição recusada para evitar apagar a base por engano.")
        # Gerente vigente = último mês disponível, independentemente da ordem das linhas.
        for uid, unidade in unis.items():
            ultimo = max(f["mes"] for f in fatos.values() if f["unidade_id"] == uid)
            unidade["gerente"] = gerentes[unidade["marca"], unidade["loja"], ultimo]
        for chave, (resumo, ctx) in resumos.items():
            marca, loja, mes = chave
            registros = [f for f in fatos.values() if f["unidade_id"] == identificador(marca, loja) and f["mes"] == mes]
            if not registros:
                raise ValueError(f"{ctx}: resumo de gerente sem consultores correspondentes.")
            for col, campo in (("Passagens", "passagens"), ("Refil Diant. (qtd)", "refil_diant"), ("Refil Tras. (qtd)", "refil_tras")):
                if resumo.get(col) is not None:
                    esperado = numero(resumo[col], f"{ctx}, {col}", inteiro=True)
                    valores = [f[campo] for f in registros]
                    if None not in valores and sum(valores) != esperado:
                        raise ValueError(f"{ctx}: {col} diverge da soma dos consultores ({sum(valores)}).")
            if any(isinstance(resumo.get(col), (dt.date, dt.datetime)) for col in ("Total Diant. (R$)", "Total Tras. (R$)")):
                avisos.append(f"{ctx}: totais formatados como datas foram ignorados; faturamento recalculado por quantidade × preço.")
        tabelas["unidades"] = list(unis.values())
        tabelas["consultores"] = list(cons.values())
        tabelas["lancamentos"] = list(fatos.values())
        for cid, itens in vinculos.items():
            ultimo = max(m for m, _ in itens)
            for uid in sorted({u for m, u in itens if m == ultimo}):
                tabelas["consultor_unidade"].append(dict(consultor_id=cid, unidade_id=uid))
        ler_verbas(wb, tabelas, avisos)
        return Base(tabelas, avisos, hashlib.sha256(conteudo).hexdigest(), caminho.name)
    finally:
        wb.close()


VERBAS_COLS = ["Data", "Pedido", "Cliente", "Produto", "Tipo", "Quantidade", "Preço Unitário",
               "Verba Consultor", "Verba Gerente", "Verba Marketing"]


def ler_verbas(wb, tabelas, avisos):
    if "Verbas" in wb.sheetnames:
        for i, (r, ctx) in enumerate(linhas(wb, "Verbas", VERBAS_COLS), 1):
            tipo = texto(r["Tipo"], ctx).lower()
            if tipo not in ("diant", "tras"):
                raise ValueError(f"{ctx}: Tipo deve ser diant ou tras.")
            n = numero(r["Quantidade"], ctx, inteiro=True)
            p = numero(r["Preço Unitário"], ctx)
            c, g, m = (numero(r[col], f"{ctx}, {col}") for col in VERBAS_COLS[-3:])
            tabelas["vendas_verbas"].append(dict(id=i, data=data(r["Data"], ctx), pedido=texto(r["Pedido"], ctx, 100),
                cliente=texto(r["Cliente"], ctx), produto=texto(r["Produto"], ctx), tipo_refil=tipo,
                qtde=n, preco_unit=p, total_item=n*p, verba_consultor=c, total_consultor=n*c,
                verba_gerente=g, total_gerente=n*g, verba_reserva=m, total_reserva=n*m))
    else:
        avisos.append("Sem aba Verbas: verbas ficarão sem dados (a carga substitui toda a base operacional).")
    meses = set()
    if "Pagamentos" in wb.sheetnames:
        for r, ctx in linhas(wb, "Pagamentos", ["Mês", "Consultor Pago", "Gerente Pago"]):
            mes = data(r["Mês"], ctx, mes=True)
            if mes in meses:
                raise ValueError(f"{ctx}: mês de pagamento duplicado.")
            meses.add(mes)
            flags = []
            for col in ("Consultor Pago", "Gerente Pago"):
                val = str(r[col]).strip().lower()
                if val not in ("sim", "não", "nao"):
                    raise ValueError(f"{ctx}: {col} deve ser Sim ou Não.")
                flags.append(val == "sim")
            tabelas["verbas_pagamentos"].append(dict(mes=mes, consultor_pago=flags[0], gerente_pago=flags[1]))
    marketing = {}
    if "Marketing" in wb.sheetnames:
        for r, ctx in linhas(wb, "Marketing", ["Mês", "Valor"]):
            mes = data(r["Mês"], ctx, mes=True)
            marketing[mes] = marketing.get(mes, Decimal(0)) + numero(r["Valor"], ctx)
    tabelas["verbas_marketing_pagos"] = [dict(mes=m, valor=v) for m, v in marketing.items()]
