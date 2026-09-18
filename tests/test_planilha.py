import datetime as dt
from decimal import Decimal

from openpyxl import load_workbook
import pytest

from configuracao import BASE_DIR
from planilha import ler_planilha, numero


def alterada(tmp_path, editar):
    wb = load_workbook(BASE_DIR / "BASE GRUPO SAGA.xlsx")
    editar(wb)
    p = tmp_path / "teste.xlsx"
    wb.save(p)
    wb.close()
    return p


def test_base_e_totais_de_gerentes_com_datas():
    base = ler_planilha(BASE_DIR / "BASE GRUPO SAGA.xlsx")
    assert base.faturamento == Decimal("2275.00")
    assert len(base.tabelas["lancamentos"]) == 6
    assert sum(r["passagens"] for r in base.tabelas["lancamentos"]) == 428
    assert {r["gerente"] for r in base.tabelas["unidades"]} == {"Gerente 1", "Gerente 2"}
    assert any("datas" in a for a in base.avisos)


@pytest.mark.parametrize("valor", [-1, 1.2, float("nan"), float("inf"), dt.date.today(), True])
def test_quantidades_invalidas(valor):
    with pytest.raises(ValueError):
        numero(valor, "Teste", inteiro=True)


def test_moeda_brasileira_e_zero():
    assert numero("R$ 1.234,56", "teste") == Decimal("1234.56")
    assert numero(0, "teste", opcional=True) == 0
    assert numero(None, "teste", opcional=True) is None


def test_duplicado_bloqueado(tmp_path):
    p = alterada(tmp_path, lambda w: w["Consultores"].append([c.value for c in w["Consultores"][2]]))
    with pytest.raises(ValueError, match="duplicado"):
        ler_planilha(p)


def test_base_vazia_bloqueada(tmp_path):
    p = alterada(tmp_path, lambda w: w["Consultores"].delete_rows(2, 99))
    with pytest.raises(ValueError, match="vazia"):
        ler_planilha(p)


def test_gerente_divergente_bloqueado(tmp_path):
    p = alterada(tmp_path, lambda w: setattr(w["Gerentes"]["F2"], "value", 999))
    with pytest.raises(ValueError, match="diverge"):
        ler_planilha(p)


def test_passagens_ausentes_nao_viram_zero(tmp_path):
    p = alterada(tmp_path, lambda w: setattr(w["Consultores"]["F2"], "value", None))
    base = ler_planilha(p)
    assert base.tabelas["lancamentos"][0]["passagens"] is None


def test_precos_por_unidade_e_historicos(tmp_path):
    def editar(w):
        w["Preço por Unidade"]["C3"] = 100
        w["Consultores"]["M5"] = None
    base = ler_planilha(alterada(tmp_path, editar))
    fatos = base.tabelas["lancamentos"]
    assert fatos[0]["preco_diant"] == 10
    assert fatos[3]["preco_diant"] == 100
    assert fatos[4]["preco_diant"] == 10


def test_verbas_pagamentos_marketing(tmp_path):
    from planilha import VERBAS_COLS
    def editar(w):
        s = w.create_sheet("Verbas")
        s.append(VERBAS_COLS)
        s.append([dt.date(2026,9,1), "Pedido 1", "Loja 1", "Refil", "diant", 2, 10, 1, 2, 3])
        s = w.create_sheet("Pagamentos")
        s.append(["Mês", "Consultor Pago", "Gerente Pago"])
        s.append(["09/2026", "Sim", "Não"])
        s = w.create_sheet("Marketing")
        s.append(["Mês", "Valor"])
        s.append(["09/2026", "R$ 1.000,00"])
        s.append(["09/2026", 250])
    base = ler_planilha(alterada(tmp_path, editar))
    assert base.tabelas["vendas_verbas"][0]["total_reserva"] == 6
    assert base.tabelas["verbas_pagamentos"][0]["consultor_pago"] is True
    assert base.tabelas["verbas_pagamentos"][0]["gerente_pago"] is False
    assert base.tabelas["verbas_marketing_pagos"][0]["valor"] == 1250
